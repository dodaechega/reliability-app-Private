# 엑셀 내용과 HTML 이스케이프를 실제 산출물로 검증한다.
from io import BytesIO

from openpyxl import load_workbook

import db
from compare import build_matrix, render_html
from export_excel import build_model_xlsx


def test_xlsx_round_trip_and_formula_text(model_id):
    items = db.get_model_items(model_id)
    items[0].update(spec='=HYPERLINK("https://example.invalid", "link")', stage="MP")
    book = load_workbook(BytesIO(build_model_xlsx(db.get_model(model_id), items)))
    sheet = book.active
    assert sheet.max_row == len(items) + 3
    assert sheet["D3"].value == "State"
    assert sheet["D4"].value == "MP"
    assert sheet["F4"].value == items[0]["spec"]
    assert sheet["F4"].data_type == "s"
    assert sheet.freeze_panes == "A4"


def test_comparison_escapes_text_in_both_themes():
    items = {"A": [{"item_name": "Drop <script>", "spec": "<10kg & safe"}],
             "B": [{"item_name": "Drop <script>", "spec": "20kg"}]}
    rows = build_matrix(items, ["A", "B"])
    assert rows[0]["diff"]
    for dark in (False, True):
        output = render_html(rows, ["A", "B"], dark=dark)
        assert "<script>" not in output
        assert "&lt;10kg &amp; safe" in output
